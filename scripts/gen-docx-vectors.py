"""Generate the vectors of copilot-docx (docxtpl 0.20 over python-docx 1.2 and lxml 6).

    /var/lib/copilot-venv/bin/python -I scripts/gen-docx-vectors.py \
        > crates/copilot-docx/src/docx_vectors.json

Every case renders a .docx template the way `app/incidents/services/reports.py` does it:
`DocxTemplate(path).render(context, jinja_env=SandboxedEnvironment(autoescape=True))`,
then `save()`. The templates are the bundled `case_report_jinja_template.docx` (read by
the Rust test from the backend's tree) and documents built here with python-docx and
edited as a zip (tags split across runs, the row/cell/paragraph/run tags, `colspan`,
`cellbg`, `vm`, `hm`, the `{%-`/`-%}` merges, comments, listing characters, smart
quotes, a picture, headers and footers, core properties, footnotes over two sections,
an indented `document.xml`, a package without core properties).

- `templates`: `{name: base64}` of the generated templates (`@bundled` names the
  backend's file);
- `cases[i]`: `template`, `vars` (with `{"$dt": iso}` for a `datetime`), and the outcome:
  `{"members": [[name, text]|[name, {"sha256": hex}]]}` in the order `save()` writes them
  (text for the XML and relationship members Python writes or rewrites, a digest for the
  others), or `{"error": {"kind": "ValueError"|"Internal", "message"}}` (the template's
  path written `<path>`);
- the clock: `datetime.now(timezone.utc)` of a default core-properties part is `CLOCK`.
"""
import base64
import datetime
import hashlib
import io
import json
import os
import struct
import sys
import tempfile
import zipfile
import zlib

import docx
import docx.opc.parts.coreprops as coreprops_module
from docx.shared import Inches
from docxtpl import DocxTemplate
from jinja2.sandbox import SandboxedEnvironment
from lxml import etree

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUNDLED = os.path.join(ROOT, "backend", "app", "incidents", "templates", "case_report_jinja_template.docx")
CLOCK = datetime.datetime(2026, 3, 4, 5, 6, 7, tzinfo=datetime.timezone.utc)


class _FixedDatetime(datetime.datetime):
    @classmethod
    def now(cls, tz=None):
        return CLOCK


_real_timezone = datetime.timezone
_real_date = datetime.date


class _dt:
    datetime = _FixedDatetime
    timezone = _real_timezone
    date = _real_date


coreprops_module.dt = _dt

TEXT_MEMBERS = ("word/document.xml", "word/footnotes.xml", "docProps/core.xml", "[Content_Types].xml")


def png(width=4, height=3):
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def saved(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def edit_zip(data, replace=None, add=None, drop=()):
    replace = replace or {}
    add = add or {}
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename in drop:
                continue
            content = src.read(info.filename)
            if info.filename in replace:
                content = replace[info.filename](content)
            dst.writestr(info.filename, content)
        for name, content in add.items():
            dst.writestr(name, content)
    return out.getvalue()


def slim(data, drop):
    """Take parts out of a package: their members, the relationships to them, their overrides."""
    drop = set(drop)
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            name = info.filename
            if name in drop:
                continue
            content = src.read(name)
            if name.endswith(".rels"):
                base = os.path.dirname(os.path.dirname(name))
                tree = etree.fromstring(content)
                for rel in list(tree):
                    target = os.path.normpath(os.path.join(base, rel.get("Target"))).lstrip("/")
                    if target in drop:
                        tree.remove(rel)
                content = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name == "[Content_Types].xml":
                tree = etree.fromstring(content)
                for override in list(tree):
                    if (override.get("PartName") or "").lstrip("/") in drop:
                        tree.remove(override)
                content = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
            # A fixed timestamp keeps the generated templates (and this file) reproducible.
            dst.writestr(zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0)), content, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


LIGHT = ("word/stylesWithEffects.xml", "docProps/thumbnail.jpeg")
BARE = LIGHT + (
    "word/styles.xml",
    "word/theme/theme1.xml",
    "word/fontTable.xml",
    "word/webSettings.xml",
    "word/numbering.xml",
    "word/settings.xml",
)


def runs(paragraph, *parts):
    for i, text in enumerate(parts):
        run = paragraph.add_run(text)
        run.bold = i % 2 == 1


def template_main():
    d = docx.Document()
    d.core_properties.title = "{{ case.name }} report"
    d.core_properties.subject = "Case {{ case.id }}"
    d.core_properties.author = "{{ case.assigned_to }}"
    d.core_properties.comments = "{{ case.assigned_to }} & {{ case.id }}"
    d.add_heading("Case {{ case.name }}", level=1)
    runs(d.add_paragraph(), "Split ", "{{ ca", "se.na", "me }}", " across runs")
    runs(d.add_paragraph(), "Quotes: {{ case.name if case.name else “none” }} ", "and ‘{{ case.id }}’")
    d.add_paragraph("{%p if case.alerts %}")
    d.add_paragraph("There are {{ case.alerts | length }} alerts")
    d.add_paragraph("{%p endif %}")
    d.add_paragraph("{#p a paragraph comment #}")
    d.add_paragraph("Listing: {{ case.description }}")
    d.add_paragraph("Before {%- if true %} merged {% endif -%} after")
    p = d.add_paragraph("Run tag: ")
    p.add_run("{{r case.name }}")
    p.add_run(" end")
    d.add_paragraph("Escaped {_{ literal }_} and {_% literal %_}")
    d.add_paragraph("Created {{ case.case_creation_time }} closed {{ case.case_closed_time }} missing {{ nothing }}")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    t.cell(0, 0).text = "{%tr for alert in case.alerts %}"
    t.cell(1, 0).text = "{{ alert.alert_name }}"
    t.cell(1, 1).text = "{{ alert.status }}"
    t.cell(1, 2).text = "{{ alert.tags | join(', ') }}"
    t.cell(2, 0).text = "{%tr endfor %}"
    t2 = d.add_table(rows=2, cols=2)
    t2.cell(0, 0).text = "{%tc for c in cols %}"
    t2.cell(0, 1).text = "{{ c }}{%tc endfor %}"
    t2.cell(1, 0).text = "{% colspan 2 %}spanned"
    t2.cell(1, 1).text = "{% cellbg 'FF0000' %}red"
    t3 = d.add_table(rows=3, cols=2)
    t3.cell(0, 0).text = "{%tr for a in case.alerts %}"
    t3.cell(1, 0).text = "{% vm %}{{ case.name }}"
    t3.cell(1, 1).text = "{{ a.alert_name }}"
    t3.cell(2, 0).text = "{%tr endfor %}"
    t4 = d.add_table(rows=1, cols=3)
    t4.cell(0, 0).text = "{%tc for c in cols %}"
    t4.cell(0, 1).text = "{% hm %}{{ c }}"
    t4.cell(0, 2).text = "{%tc endfor %}"
    d.add_picture(io.BytesIO(png()), width=Inches(0.5))
    d.add_picture(io.BytesIO(png(2, 2)), width=Inches(0.25))
    d.add_paragraph("Grüße — {{ case.name }} & <done>")
    section = d.sections[0]
    section.header.paragraphs[0].text = "Header {{ case.name }} {{ case.id }}"
    section.footer.paragraphs[0].text = "Footer {{ case.assigned_to }}"
    d.add_section()
    d.add_paragraph("Second section {{ case.id }}")
    return saved(d)


FOOTNOTES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:id="1"><w:p><w:r><w:t>Note {{ case.name }} {{ "{{" }}x{{ "}}" }}</w:t></w:r></w:p></w:footnote>'
    "</w:footnotes>"
)


def template_footnotes():
    d = docx.Document()
    d.add_paragraph("Body {{ case.name }}")
    d.add_section()
    d.add_paragraph("Again")
    data = saved(d)

    def rels(content):
        return content.replace(
            b"</Relationships>",
            b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
        )

    def types(content):
        return content.replace(
            b"</Types>",
            b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
        )

    return edit_zip(
        data,
        replace={"word/_rels/document.xml.rels": rels, "[Content_Types].xml": types},
        add={"word/footnotes.xml": FOOTNOTES.encode()},
    )


def template_indented():
    d = docx.Document()
    d.add_paragraph("Indented {{ case.name }}")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "{{ case.id }}"
    data = saved(d)

    def pretty(content):
        tree = etree.fromstring(content)
        return etree.tostring(tree, pretty_print=True, xml_declaration=True, encoding="UTF-8", standalone=True)

    return edit_zip(data, replace={"word/document.xml": pretty, "word/styles.xml": pretty})


def template_no_core():
    d = docx.Document()
    d.add_paragraph("No core {{ case.name }}")
    data = saved(d)

    def rels(content):
        tree = etree.fromstring(content)
        for r in list(tree):
            if r.get("Target") == "docProps/core.xml":
                tree.remove(r)
        return etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)

    def types(content):
        return content.replace(
            b'<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>',
            b"",
        )

    return edit_zip(data, replace={"_rels/.rels": rels, "[Content_Types].xml": types}, drop=("docProps/core.xml",))


def template_long_title():
    d = docx.Document()
    d.core_properties.title = "{{ case.description * 30 }}"
    d.add_paragraph("x")
    return saved(d)


def template_not_word():
    data = template_no_core()

    def types(content):
        return content.replace(b"document.main+xml", b"template.main+xml")

    return edit_zip(data, replace={"[Content_Types].xml": types})


def template_syntax_error():
    d = docx.Document()
    d.add_paragraph("{% for x in %}")
    return saved(d)


def template_value_error():
    d = docx.Document()
    d.add_paragraph("{{ range(1, 2, 0) }}")
    return saved(d)


def template_control_chars():
    d = docx.Document()
    d.add_paragraph("Value: {{ case.description }}")
    return saved(d)


TEMPLATES = {
    "main": slim(template_main(), LIGHT),
    "footnotes": slim(template_footnotes(), LIGHT),
    "indented": slim(template_indented(), LIGHT),
    "no_core": slim(template_no_core(), BARE),
    "long_title": slim(template_long_title(), BARE),
    "not_word": slim(template_not_word(), BARE),
    "syntax_error": slim(template_syntax_error(), BARE),
    "value_error": slim(template_value_error(), BARE),
    "control_chars": slim(template_control_chars(), BARE),
}


def dt(value):
    return {"$dt": value.isoformat()}


def decode(value):
    if isinstance(value, dict):
        if "$dt" in value:
            return datetime.datetime.fromisoformat(value["$dt"])
        return {k: decode(v) for k, v in value.items()}
    if isinstance(value, list):
        return [decode(v) for v in value]
    return value


ALERT = {
    "alert_name": "Suspicious <login> & more",
    "alert_description": "Line one\nLine two\ttabbed",
    "status": "OPEN",
    "time_closed": None,
    "tags": ["wazuh", "o'brien"],
    "assets": [{"asset_name": "web-01", "agent_id": "001"}],
    "comments": [{"comment": "checked", "user_name": "admin", "created_at": dt(datetime.datetime(2026, 1, 2, 3, 4, 5))}],
    "context": {"source": "wazuh", "context": {"rule": {"id": 5710, "level": 5.5}, "ok": True, "none": None, "list": [1, "a"]}},
    "iocs": [{"ioc_value": "1.2.3.4", "ioc_type": "IP", "ioc_description": None}],
}

CASE = {
    "name": "Case “Alpha” & co",
    "description": "First\nSecond\tTab\aBell\fPage",
    "assigned_to": "analyst",
    "case_creation_time": dt(datetime.datetime(2026, 2, 3, 4, 5, 6)),
    "case_closed_time": None,
    "id": 42,
    "alerts": [ALERT, dict(ALERT, alert_name="Second", assets=[], context=None, tags=[])],
}

EMPTY_CASE = dict(CASE, alerts=[], description="", name="", assigned_to=None, case_closed_time=dt(datetime.datetime(2026, 2, 4, 0, 0, 0)))

CASES = [
    ("@bundled", {"case": CASE}),
    ("@bundled", {"case": EMPTY_CASE}),
    ("@bundled", {"case": dict(CASE, alerts=[dict(ALERT, context={"source": None, "context": None})])}),
    ("main", {"case": CASE, "cols": ["c1", "c2", "c3", "c4"]}),
    ("main", {"case": EMPTY_CASE, "cols": []}),
    ("main", {"case": dict(CASE, alerts=CASE["alerts"] * 3), "cols": ["only"]}),
    ("footnotes", {"case": CASE}),
    ("indented", {"case": CASE}),
    ("no_core", {"case": CASE}),
    ("long_title", {"case": CASE}),
    ("not_word", {"case": CASE}),
    ("syntax_error", {"case": CASE}),
    ("value_error", {"case": CASE}),
    ("control_chars", {"case": dict(CASE, description="esc\x1b[0m bell\x0b vt\x00nul")}),
]


def member(name, content):
    if name in TEXT_MEMBERS or name.endswith(".rels") or name.startswith(("word/header", "word/footer")):
        return [name, content.decode("utf-8")]
    return [name, {"sha256": hashlib.sha256(content).hexdigest()}]


def run(template_bytes, context):
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "template.docx")
        with open(path, "wb") as fh:
            fh.write(template_bytes)
        try:
            doc = DocxTemplate(path)
            doc.render(context, jinja_env=SandboxedEnvironment(autoescape=True))
            out = io.BytesIO()
            doc.save(out)
        except ValueError as exc:
            return {"error": {"kind": "ValueError", "message": str(exc).replace(path, "<path>")}}
        except Exception as exc:  # noqa: BLE001 - every other failure is a bare 500
            return {"error": {"kind": "Internal", "message": f"{type(exc).__name__}: {exc}"}}
        z = zipfile.ZipFile(io.BytesIO(out.getvalue()))
        return {"members": [member(i.filename, z.read(i.filename)) for i in z.infolist()]}


def main():
    with open(BUNDLED, "rb") as fh:
        bundled = fh.read()
    cases = []
    for name, context in CASES:
        data = bundled if name == "@bundled" else TEMPLATES[name]
        cases.append({"template": name, "vars": context, "outcome": run(data, decode(context))})
    doc = {
        "clock": CLOCK.isoformat(),
        "templates": {k: base64.b64encode(v).decode() for k, v in TEMPLATES.items()},
        "cases": cases,
    }
    json.dump(doc, sys.stdout, indent=1, ensure_ascii=False)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
